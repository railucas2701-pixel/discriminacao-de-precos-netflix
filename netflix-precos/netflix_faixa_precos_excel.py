"""
Gera netflix_faixa_precos_regiao.xlsx: dados por país, resumo por região (com fórmulas) e o gráfico
de faixa de preço (barras empilhadas com a 1ª série invisível = "barras flutuantes").
Entrada: resultado_netflix/netflix_plano_basico_por_pais.csv e <cache>/_bm_paises.json
Uso: python netflix_faixa_precos_excel.py [--cache netflix_cache]
"""
import argparse
import csv
import json
import os

import country_converter as coco
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.formula import ArrayFormula

SAIDA = "netflix_faixa_precos_regiao.xlsx"
REGIOES = ["UCAN", "LATAM", "Europa", "Oriente Médio", "África", "APAC"]
DESCRICAO = {
    "UCAN": "EUA e Canadá", "LATAM": "América Latina e Caribe",
    "Europa": "Inclui Turquia, Chipre e Cáucaso", "Oriente Médio": "Países árabes da Ásia, Israel e Irã",
    "África": "Continente inteiro", "APAC": "Ásia-Pacífico",
}
# Texto descritivo (não se atualiza sozinho se os preços mudarem)
PAISES_MIN = {"UCAN": "Canadá · Standard c/ anúncios", "LATAM": "Belize e mais 8 países · Básico",
              "Europa": "Turquia · Básico", "Oriente Médio": "Iraque, Irã e mais 4 países · Básico",
              "África": "Quênia · Mobile", "APAC": "Paquistão · Mobile"}
PAISES_MAX = {"UCAN": "EUA · Standard c/ anúncios", "LATAM": "Uruguai e Trinidad e Tobago · Básico",
              "Europa": "Suíça e Liechtenstein · Básico", "Oriente Médio": "Israel · Básico",
              "África": "Marrocos · Básico", "APAC": "Singapura · Básico"}

FONTE = "Arial"
AZUL_INPUT = Font(name=FONTE, color="0000FF")
PRETO = Font(name=FONTE)
NEGRITO = Font(name=FONTE, bold=True)
CINZA = Font(name=FONTE, color="595959", italic=True, size=9)
CAB = Font(name=FONTE, bold=True, color="FFFFFF")
FUNDO_CAB = PatternFill("solid", fgColor="1F3864")
LINHA = Border(bottom=Side(style="thin", color="D9D9D9"))


def regiao6(p, wb, continente):
    if p["cc"] == "IR":
        return "Oriente Médio"
    if p["regiao"] != "EMEA":
        return p["regiao"]
    if continente == "Africa":
        return "África"
    if continente == "Europe":
        return "Europa"
    return "Europa" if wb.get(p["cc"]) == "Europe & Central Asia" else "Oriente Médio"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache", default="netflix_cache")
    args = ap.parse_args()
    paises = [p for p in csv.DictReader(open("resultado_netflix/netflix_plano_basico_por_pais.csv", encoding="utf-8"))
              if p["tipo"] == "país"]
    _, bm = json.load(open(os.path.join(args.cache, "_bm_paises.json")))
    wb_reg = {p["iso2Code"]: p["region"]["value"].strip() for p in bm}
    continentes = coco.convert(names=[p["cc"] for p in paises], src="ISO2", to="continent")
    for p, c in zip(paises, continentes):
        p["reg6"] = regiao6(p, wb_reg, c)
    paises.sort(key=lambda p: (REGIOES.index(p["reg6"]), p["pais"]))

    wb = Workbook()
    # ---------------------------------------------------------------- aba Países
    ws_p = wb.active
    ws_p.title = "Países"
    cab = ["País", "Região", "Plano mais básico", "Preço (US$/mês)", "Moeda de cobrança", "Preço local"]
    ws_p.append(cab)
    for c in ws_p[1]:
        c.font, c.fill = CAB, FUNDO_CAB
    for p in paises:
        ws_p.append([p["pais"], p["reg6"], p["plano_mais_basico"], float(p["preco_usd"]), p["moeda"],
                     float(p["preco_local"])])
    n = len(paises)
    for linha in ws_p.iter_rows(min_row=2, max_row=n + 1):
        for c in linha:
            c.font = AZUL_INPUT
        linha[3].number_format = '"US$" #,##0.00'
        linha[5].number_format = "#,##0.00"
    for col, w in zip("ABCDEF", (34, 16, 22, 16, 18, 14)):
        ws_p.column_dimensions[col].width = w
    ws_p.freeze_panes = "A2"
    ws_p.cell(row=n + 3, column=1, value=(
        "Fonte: Central de Ajuda da Netflix (help.netflix.com/en/node/24926/<país>), coletado em 10/10/2026; "
        "câmbio do mesmo dia (open.er-api.com). Só países (membros da ONU + Vaticano, Palestina e Taiwan); "
        "territórios excluídos. Valores em azul são dados coletados.")).font = CINZA
    ws_p.cell(row=n + 4, column=1, value=(
        "Regiões: África = continente inteiro; Europa = continente europeu + Turquia, Chipre e Cáucaso "
        "(classificação do Banco Mundial); Oriente Médio = países árabes da Ásia, Israel e Irã; demais pela "
        "geografia da ONU. A Netflix não publica a lista de países de cada região.")).font = CINZA

    # ---------------------------------------------------------------- aba Faixa de preço
    ws = wb.create_sheet("Faixa de preço", 0)
    ws["A1"] = "Netflix: faixa de preço do plano mais básico por região (US$/mês, outubro de 2026)"
    ws["A1"].font = Font(name=FONTE, bold=True, size=14)
    ws["A2"] = "Mínimo, mediana e máximo calculados por fórmula a partir da aba Países."
    ws["A2"].font = CINZA

    cab = ["Região", "Descrição", "Países", "Mínimo (US$)", "Mediana (US$)", "Máximo (US$)",
           "País(es) no mínimo", "País(es) no máximo",
           "Base do gráfico (invisível)", "Metade mais barata (mín. → mediana)", "Metade mais cara (mediana → máx.)"]
    ws.append([])
    ws.append(cab)  # linha 4
    for c in ws[4]:
        c.font, c.fill = CAB, FUNDO_CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 32

    reg = f"'Países'!$B$2:$B${n + 1}"
    preco = f"'Países'!$D$2:$D${n + 1}"
    for i, r in enumerate(REGIOES):
        lin = 5 + i
        ws.cell(row=lin, column=1, value=r).font = NEGRITO
        ws.cell(row=lin, column=2, value=DESCRICAO[r]).font = PRETO
        ws.cell(row=lin, column=3, value=f"=COUNTIF({reg},A{lin})")
        ws.cell(row=lin, column=4, value=f"=_xlfn.MINIFS({preco},{reg},A{lin})")
        ws[f"E{lin}"] = ArrayFormula(f"E{lin}", f"=MEDIAN(IF({reg}=A{lin},{preco}))")
        ws.cell(row=lin, column=6, value=f"=_xlfn.MAXIFS({preco},{reg},A{lin})")
        ws.cell(row=lin, column=7, value=PAISES_MIN[r]).font = AZUL_INPUT
        ws.cell(row=lin, column=8, value=PAISES_MAX[r]).font = AZUL_INPUT
        ws.cell(row=lin, column=9, value=f"=D{lin}")
        ws.cell(row=lin, column=10, value=f"=E{lin}-D{lin}")
        ws.cell(row=lin, column=11, value=f"=F{lin}-E{lin}")
        for col in range(1, 12):
            c = ws.cell(row=lin, column=col)
            c.border = LINHA
            if col in (3, 4, 5, 6, 9, 10, 11) and c.font != NEGRITO:
                c.font = PRETO
            if col in (4, 5, 6, 9, 10, 11):
                c.number_format = '"US$" #,##0.00'
    ult = 4 + len(REGIOES)
    ws.cell(row=ult + 1, column=7, value=(
        "Colunas G e H (azul) são texto fixo: não mudam se os preços da aba Países mudarem.")).font = CINZA
    ws.cell(row=ult + 2, column=1, value=(
        "Colunas I a K alimentam o gráfico: a 1ª série (base) fica sem preenchimento, então cada barra começa "
        "no mínimo e termina no máximo; a troca de cor marca a mediana.")).font = CINZA
    for col, w in zip("ABCDEFGHIJK", (15, 32, 8, 13, 14, 13, 34, 36, 13, 15, 15)):
        ws.column_dimensions[col].width = w

    # Gráfico: barras horizontais empilhadas com a 1ª série invisível
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 80
    ch.title = "Faixa de preço do plano mais básico (US$/mês)"
    dados = Reference(ws, min_col=9, max_col=11, min_row=4, max_row=ult)
    cats = Reference(ws, min_col=1, min_row=5, max_row=ult)
    ch.add_data(dados, titles_from_data=True)
    ch.set_categories(cats)
    base, baixo, alto = ch.series
    base.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    baixo.graphicalProperties = GraphicalProperties(solidFill="86B6EF", ln=LineProperties(noFill=True))
    alto.graphicalProperties = GraphicalProperties(solidFill="2A78D6", ln=LineProperties(noFill=True))
    ch.x_axis.scaling.orientation = "maxMin"  # 1ª região no topo, como na tabela
    ch.y_axis.crosses = "max"                 # mantém o eixo de valores embaixo
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 20
    ch.y_axis.majorUnit = 5
    ch.y_axis.number_format = '"US$" 0'
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9"))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    ch.legend.legendEntry = [LegendEntry(idx=0, delete=True)]  # esconde a série invisível
    ch.height, ch.width = 9.5, 18
    ws.add_chart(ch, f"A{ult + 4}")

    # ---------------------------------------------------------------- aba Como fazer
    wi = wb.create_sheet("Como fazer")
    passos = [
        "Como montar o gráfico de faixa (mínimo a máximo) no Excel",
        "",
        "Opção 1 - Barras flutuantes (a usada na aba 'Faixa de preço'):",
        "1. Monte 3 colunas: Base = mínimo; Parte 1 = mediana - mínimo; Parte 2 = máximo - mediana.",
        "   (sem mediana, use só 2 colunas: Base = mínimo e Faixa = máximo - mínimo)",
        "2. Selecione as regiões e as 3 colunas > Inserir > Gráfico de Barras > Barras Empilhadas.",
        "3. Clique na série 'Base' > Formatar Série de Dados > Preenchimento: Sem preenchimento; Borda: Sem linha.",
        "4. Para a 1ª região ficar no topo: Formatar eixo das categorias > 'Categorias em ordem inversa'",
        "   e 'Eixo horizontal cruza: na categoria máxima'.",
        "5. Fixe o eixo de valores (ex.: mínimo 0, máximo 20) e apague da legenda a entrada 'Base'.",
        "",
        "Opção 2 - Gráfico de Ações 'Máxima-Mínima-Fechamento' (nome pode variar na sua versão do Excel):",
        "- Precisa de 3 colunas nesta ordem: Máximo, Mínimo, Fechamento (use a mediana como 'fechamento').",
        "- Inserir > Gráficos de Cascata, Funil, Ações... > Ações > Máxima-Mínima-Fechamento.",
        "- Desenha uma linha vertical do mínimo ao máximo por região, com um tique na mediana.",
        "- Fica vertical e é mais difícil de formatar; para slide, as barras flutuantes costumam ficar melhores.",
        "",
        "Opção 3 - Gráfico de linhas com 'Linhas de máximo/mínimo':",
        "- Faça um gráfico de linhas com as séries Mínimo e Máximo, depois Adicionar Elemento do Gráfico >",
        "  Linhas > Linhas de Máximo/Mínimo, e deixe as linhas das séries sem cor (fica só o traço vertical).",
        "",
        "Para levar ao PowerPoint: copie o gráfico e cole no slide (Colar > Manter formatação de origem e",
        "inserir pasta de trabalho, para continuar editável).",
    ]
    for i, t in enumerate(passos, 1):
        c = wi.cell(row=i, column=1, value=t)
        c.font = Font(name=FONTE, bold=(i == 1 or t.startswith("Opção")), size=13 if i == 1 else 11)
    wi.column_dimensions["A"].width = 110

    wb.save(SAIDA)
    print("ok:", SAIDA, "| países:", n)


if __name__ == "__main__":
    main()
